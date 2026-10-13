from pathlib import Path
from random import Random

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


OUTPUT = Path("templates/RFM_Contact_Heatmap_Template.xlsx")
MAX_DATA_ROW = 100005

NAVY = "17365D"
BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "EAF2F8"
WHITE = "FFFFFF"
GRAY = "666666"
LIGHT_GRAY = "E7E6E6"
RED = "F8696B"
YELLOW = "FFEB84"
GREEN = "63BE7B"


def style_title(ws, cell_range, text):
    ws.merge_cells(cell_range)
    cell = ws[cell_range.split(":")[0]]
    cell.value = text
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.font = Font(color=WHITE, bold=True, size=16)
    cell.alignment = Alignment(horizontal="left", vertical="center")


def style_section(ws, cell_range, text):
    ws.merge_cells(cell_range)
    cell = ws[cell_range.split(":")[0]]
    cell.value = text
    cell.fill = PatternFill("solid", fgColor=BLUE)
    cell.font = Font(color=WHITE, bold=True, size=11)
    cell.alignment = Alignment(horizontal="left", vertical="center")


def add_heatmap(ws, top_row, title, formula_kind, number_format):
    style_section(ws, f"A{top_row}:G{top_row}", title)
    ws.cell(top_row + 1, 1, "R \\ F")
    for col, score in enumerate(range(1, 6), start=2):
        ws.cell(top_row + 1, col, score)
    ws.cell(top_row + 1, 7, "Total")

    for offset, r_score in enumerate(range(1, 6), start=2):
        row = top_row + offset
        ws.cell(row, 1, r_score)
        for col in range(2, 7):
            f_cell = ws.cell(top_row + 1, col).coordinate
            r_cell = ws.cell(row, 1).coordinate
            if formula_kind == "sum":
                formula = (
                    f'=SUMIFS(\'Datos Contacto\'!$G$6:$G${MAX_DATA_ROW},'
                    f'\'Datos Contacto\'!$E$6:$E${MAX_DATA_ROW},${r_cell},'
                    f'\'Datos Contacto\'!$F$6:$F${MAX_DATA_ROW},{f_cell},'
                    f'\'Datos Contacto\'!$C$6:$C${MAX_DATA_ROW},"Qualified")'
                )
            elif formula_kind == "average":
                formula = (
                    f'=IFERROR(AVERAGEIFS(\'Datos Contacto\'!$G$6:$G${MAX_DATA_ROW},'
                    f'\'Datos Contacto\'!$E$6:$E${MAX_DATA_ROW},${r_cell},'
                    f'\'Datos Contacto\'!$F$6:$F${MAX_DATA_ROW},{f_cell},'
                    f'\'Datos Contacto\'!$C$6:$C${MAX_DATA_ROW},"Qualified"),0)'
                )
            else:
                formula = (
                    f'=COUNTIFS(\'Datos Contacto\'!$E$6:$E${MAX_DATA_ROW},${r_cell},'
                    f'\'Datos Contacto\'!$F$6:$F${MAX_DATA_ROW},{f_cell},'
                    f'\'Datos Contacto\'!$C$6:$C${MAX_DATA_ROW},"Qualified")'
                )
            ws.cell(row, col, formula)
            ws.cell(row, col).number_format = number_format
        ws.cell(row, 7, f"=SUM(B{row}:F{row})")
        ws.cell(row, 7).number_format = number_format

    total_row = top_row + 7
    ws.cell(total_row, 1, "Total")
    for col in range(2, 8):
        letter = get_column_letter(col)
        ws.cell(total_row, col, f"=SUM({letter}{top_row + 2}:{letter}{top_row + 6})")
        ws.cell(total_row, col).number_format = number_format

    header = ws[f"A{top_row + 1}:G{top_row + 1}"][0]
    for cell in header:
        cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        cell.font = Font(bold=True, color=NAVY)
        cell.alignment = Alignment(horizontal="center")
    for row in range(top_row + 2, top_row + 7):
        ws.cell(row, 1).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        ws.cell(row, 1).font = Font(bold=True, color=NAVY)
        ws.cell(row, 1).alignment = Alignment(horizontal="center")
    for cell in ws[total_row]:
        if cell.column <= 7:
            cell.fill = PatternFill("solid", fgColor=PALE_BLUE)
            cell.font = Font(bold=True, color=NAVY)

    grid = f"B{top_row + 2}:F{top_row + 6}"
    ws.conditional_formatting.add(
        grid,
        ColorScaleRule(
            start_type="min", start_color=RED,
            mid_type="percentile", mid_value=50, mid_color=YELLOW,
            end_type="max", end_color=GREEN,
        ),
    )
    thin = Side(style="thin", color="B7C9D6")
    for row in ws.iter_rows(min_row=top_row + 1, max_row=total_row, min_col=1, max_col=7):
        for cell in row:
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(horizontal="center", vertical="center")


def build():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Heatmap RFM"
    data = wb.create_sheet("Datos Contacto")
    guide = wb.create_sheet("Guía")

    # Heatmap dashboard
    ws.sheet_view.showGridLines = False
    style_title(ws, "A1:G1", "Heatmap RFM — Contactos")
    ws.row_dimensions[1].height = 28
    ws["A2"] = "Ejes: Recency (filas) × Frequency (columnas). El color representa Monetary dentro de cada cohorte."
    ws.merge_cells("A2:G2")
    ws["A2"].font = Font(italic=True, color=GRAY)
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws["A3"] = "Rojo = menor valor | Amarillo = intermedio | Verde = mayor valor"
    ws.merge_cells("A3:G3")
    ws["A3"].font = Font(color=GRAY, size=10)

    add_heatmap(ws, 5, "Principal — suma de Monetary Score por cohorte", "sum", "0")
    add_heatmap(ws, 15, "Diagnóstico — promedio de Monetary Score por cohorte", "average", "0.00")
    add_heatmap(ws, 25, "Volumen — número de contactos por cohorte", "count", "0")

    ws["A35"] = "Conciliación de registros calificados"
    ws["A35"].font = Font(bold=True, color=NAVY)
    ws["A36"] = "Registros esperados según el query mostrado"
    ws["B36"] = 48003
    ws["A37"] = "Registros Qualified importados"
    ws["B37"] = f'=COUNTIF(\'Datos Contacto\'!$C$6:$C${MAX_DATA_ROW},"Qualified")'
    ws["A38"] = "Incluidos en heatmap (R/F 1–5)"
    ws["B38"] = f'=COUNTIFS(\'Datos Contacto\'!$C$6:$C${MAX_DATA_ROW},"Qualified",\'Datos Contacto\'!$E$6:$E${MAX_DATA_ROW},">=1",\'Datos Contacto\'!$E$6:$E${MAX_DATA_ROW},"<=5",\'Datos Contacto\'!$F$6:$F${MAX_DATA_ROW},">=1",\'Datos Contacto\'!$F$6:$F${MAX_DATA_ROW},"<=5")'
    ws["A39"] = "Excluidos del heatmap por R=0 o F=0"
    ws["B39"] = "=B37-B38"
    ws["A40"] = "Cobertura del heatmap sobre Qualified"
    ws["B40"] = "=IFERROR(B38/B37,0)"
    ws["B40"].number_format = "0.0%"
    ws["A41"] = "Diferencia contra registros esperados"
    ws["B41"] = "=B37-B36"
    for row in range(35, 42):
        ws.cell(row, 1).fill = PatternFill("solid", fgColor=PALE_BLUE)
    ws.column_dimensions["A"].width = 36
    for col in "BCDEF":
        ws.column_dimensions[col].width = 14
    ws.column_dimensions["G"].width = 15
    ws.freeze_panes = "B6"

    # Editable input sheet
    data.sheet_view.showGridLines = False
    style_title(data, "A1:H1", "Datos de Contacto para el heatmap")
    data["A2"] = "Pega en A5:H... la salida del query de Salesforce, incluidos sus encabezados API. Sustituye primero las filas de ejemplo."
    data.merge_cells("A2:H2")
    data["A2"].font = Font(italic=True, color=GRAY)
    data["A3"] = "Fuente funcional: docs/RFM-Contact/FIELD_REFERENCE.md"
    data.merge_cells("A3:H3")
    data["A3"].font = Font(color=GRAY, size=10)
    headers = [
        "Id", "Name", "RFM_Qualification_Status__c", "RFM_Donor_Type__c",
        "RFM_Recency_Score__c", "RFM_Frequency_Score__c", "RFM_Monetary_Score__c", "RFM_Final_Score__c",
    ]
    for col, header in enumerate(headers, 1):
        data.cell(5, col, header)

    rng = Random(20260821)
    donor_types = ["Recurring Only", "One-Time Only", "Mixed"]
    rows = []
    idx = 1
    for r_score in range(1, 6):
        for f_score in range(1, 6):
            repetitions = 1 + ((r_score + f_score) % 3)
            for rep in range(repetitions):
                center = round((r_score + f_score) / 2)
                m_score = max(1, min(5, center + rng.choice([-1, 0, 0, 1])))
                rows.append([
                    f"003EXAMPLE{idx:05d}", f"Contacto ejemplo {idx:02d}", "Qualified",
                    donor_types[(idx - 1) % len(donor_types)], r_score, f_score, m_score,
                    f"=ROUND((0.40*E{idx + 5})+(0.35*F{idx + 5})+(0.25*G{idx + 5}),2)",
                ])
                idx += 1
    special_row = 6 + len(rows)
    rows.extend([
        ["003000000363804", "Qualified con Frequency cero", "Qualified", "Mixed", 1, 0, 2, f"=ROUND((0.40*E{special_row})+(0.35*F{special_row})+(0.25*G{special_row}),2)"],
        ["003000000198653", "Qualified con Monetary cero", "Qualified", "One-Time Only", 3, 1, 0, f"=ROUND((0.40*E{special_row + 1})+(0.35*F{special_row + 1})+(0.25*G{special_row + 1}),2)"],
    ])
    for row_idx, row in enumerate(rows, start=6):
        for col_idx, value in enumerate(row, start=1):
            data.cell(row_idx, col_idx, value)

    last_row = 5 + len(rows)
    data.freeze_panes = "A6"
    widths = [19, 25, 27, 22, 16, 17, 16, 14]
    for col, width in enumerate(widths, start=1):
        data.column_dimensions[get_column_letter(col)].width = width
    data.column_dimensions["H"].width = 14
    for cell in data[5]:
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for row in range(6, last_row + 1):
        for col in range(5, 9):
            data.cell(row, col).number_format = "0.00" if col == 8 else "0"
            data.cell(row, col).alignment = Alignment(horizontal="center")

    # Guide and assumptions
    guide.sheet_view.showGridLines = False
    style_title(guide, "A1:D1", "Guía de uso y definiciones")
    guide_rows = [
        ("Objetivo", "Explorar cohortes de Contact por scores R y F, usando Monetary como tercera dimensión de color."),
        ("Matriz principal", "Suma de RFM_Monetary_Score__c por combinación R×F. Mide valor monetario acumulado, pero también está influida por el tamaño de la cohorte."),
        ("Promedio", "Promedio de RFM_Monetary_Score__c por cohorte. Facilita comparar la calidad monetaria sin el efecto directo del volumen."),
        ("Volumen", "Conteo de contactos calificados por cohorte. Debe revisarse junto con la matriz principal."),
        ("Recency", "Contact.RFM_Recency_Score__c, de 0 a 5; 5 es mejor/más reciente."),
        ("Frequency", "Contact.RFM_Frequency_Score__c, de 0 a 5; 5 es mejor/más frecuente."),
        ("Monetary", "Contact.RFM_Monetary_Score__c, de 0 a 5; proviene del percentil monetario balanceado reciente/histórico."),
        ("Exclusión", "La matriz 5×5 incluye RFM_Qualification_Status__c = Qualified y scores R/F de 1 a 5. Los Qualified con R=0 o F=0 se concilian por separado."),
        ("Actualización", "Pega la salida del query en A5:H..., incluidos los encabezados API. Las fórmulas leen hasta la fila 100005, suficiente para más de 48,003 registros."),
        ("Query esperado", "SELECT Id, Name, RFM_Qualification_Status__c, RFM_Donor_Type__c, RFM_Recency_Score__c, RFM_Frequency_Score__c, RFM_Monetary_Score__c, RFM_Final_Score__c FROM Contact WHERE RFM_Qualification_Status__c = 'Qualified'"),
        ("Datos incluidos", "Las filas entregadas son ejemplos ficticios para visualizar el comportamiento; deben reemplazarse con una exportación real."),
    ]
    guide["A3"] = "Concepto"
    guide["B3"] = "Definición"
    for cell in guide[3][:2]:
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(color=WHITE, bold=True)
    for row_idx, (concept, definition) in enumerate(guide_rows, start=4):
        guide.cell(row_idx, 1, concept).font = Font(bold=True, color=NAVY)
        guide.cell(row_idx, 2, definition)
        guide.cell(row_idx, 2).alignment = Alignment(wrap_text=True, vertical="top")
        guide.row_dimensions[row_idx].height = 34
    guide.column_dimensions["A"].width = 24
    guide.column_dimensions["B"].width = 95
    guide.freeze_panes = "A4"

    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(OUTPUT)

    # Structural validation after re-opening.
    check = load_workbook(OUTPUT, data_only=False)
    assert check.sheetnames == ["Heatmap RFM", "Datos Contacto", "Guía"]
    assert check["Heatmap RFM"]["B7"].value.startswith("=SUMIFS")
    assert [check["Heatmap RFM"].cell(row, 1).value for row in range(7, 12)] == [1, 2, 3, 4, 5]
    assert [check["Datos Contacto"].cell(5, col).value for col in range(1, 9)] == headers
    assert "$100005" in check["Heatmap RFM"]["B7"].value
    assert check["Heatmap RFM"]["B36"].value == 48003
    assert check["Heatmap RFM"]["B39"].value == "=B37-B38"
    assert len(check["Heatmap RFM"].conditional_formatting) == 3
    assert len(check["Datos Contacto"].tables) == 0
    assert check["Datos Contacto"].auto_filter.ref is None
    assert check["Heatmap RFM"].auto_filter.ref is None
    assert check["Datos Contacto"].max_row == last_row
    formulas = [cell.value for row in check["Heatmap RFM"].iter_rows() for cell in row if isinstance(cell.value, str) and cell.value.startswith("=")]
    assert formulas and not any("#REF!" in formula for formula in formulas)
    check.close()
    print(f"CREATED={OUTPUT.resolve()}")
    print(f"SHEETS=3 SAMPLE_ROWS={len(rows)} FORMULAS={len(formulas)} CF_RULES=3")


if __name__ == "__main__":
    build()
